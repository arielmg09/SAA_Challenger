"""
SAA Challenger Optimiser
========================

Independent, simplified re-implementation of the Bank's @RISK / RISKOptimizer SAA
optimisation, built for model validation purposes (understanding the mechanics and
benchmarking the outputs). It is NOT a replacement for the production model.

Two layers are run for every strategy:

  1. REPLICA   - mirrors the RISKOptimizer set-up described in the SAA Instructions:
                 * decision variables: 7 integers in [0, 1000], normalised to weights
                 * objective: maximise mean of the portfolio's gross value (capital &
                   income) at the end of the horizon (assumed to be W97, year 20)
                 * soft constraints with quadratic penalties:
                     (1000 * deviation)^2 for Annualised Volatility and Min Yield
                     (100  * deviation)^2 for all other constraints
                 * Monte Carlo with rank-order (Spearman) correlation induced via the
                   Iman-Conover method (the approach @RISK uses)
                 * genetic-style global search (scipy differential evolution) run in two
                   passes: 500 iterations, then 1,000 iterations seeded from pass 1
  2. ANALYTICAL - deterministic benchmark: maximise w'mu subject to HARD constraints
                 (bounds, sum = 1, sqrt(w' Sigma w) <= target, w'y >= min yield).
                 Because annual returns are simulated i.i.d. with fixed weights,
                 E[terminal value] = V0 * (1 + w'mu)^H, which is monotone in w'mu.
                 The analytical optimum is therefore the answer the replica is
                 converging towards, absent simulation noise, penalties and rounding.

Usage
-----
  python saa_challenger.py --make-template SAA_Challenger_Inputs.xlsx
  python saa_challenger.py --input SAA_Challenger_Inputs.xlsx --output SAA_Challenger_Results.xlsx
  python saa_challenger.py                      # runs on embedded default values
  python saa_challenger.py --input X.xlsx --strategies Strategy_3 Strategy_6

Key assumptions (to be confirmed with the model owner)
-----------------------------------------------------
  A1  W97 = mean gross portfolio value at year 20, starting value V0 (Settings).
  A2  Expected returns are arithmetic annual means; returns are i.i.d. across years;
      portfolio rebalanced to target weights annually; income reinvested.
  A3  Marginal distribution is user-selectable (Normal or Lognormal on 1+r).
  A4  'Annualised Volatility' is measured as the standard deviation of simulated
      annual portfolio returns (replica) / sqrt(w' Sigma w) (analytical), and is
      treated as an upper limit unless VolConstraintType = 'target'.
  A5  Penalty deviations are measured in decimal units (1% = 0.01) and subtracted
      from the objective in the same units as the portfolio value.
  A6  Portfolio yield = w'y (linear).
"""

from __future__ import annotations

import argparse
import sys
import time
import warnings
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
from scipy.optimize import differential_evolution, minimize

# =====================================================================================
# 1. EMBEDDED DEFAULTS
# =====================================================================================
# Composite-level inputs below are ILLUSTRATIVE and taken from the 2022 developer
# documentation (as at 30/11/2021). Volatilities are mapped as per the 'Expected Annual
# Returns & Volatility' chart, as the printed table appears shifted by one row.
# Replace them with the current-year values via the Excel input file.

ASSETS = ["Equity", "DMGovt", "BondOther", "Property", "Gold", "HFs", "Cash"]
ALL_VARS = ASSETS + ["Inflation"]  # Inflation is simulated but not investable

ASSET_DESCRIPTIONS = {
    "Equity": "Equity",
    "DMGovt": "Bonds - developed market government",
    "BondOther": "Bonds - other",
    "Property": "Alternatives - property",
    "Gold": "Alternatives - gold",
    "HFs": "Alternatives - hedge funds",
    "Cash": "Cash",
    "Inflation": "Inflation",
}

DEFAULT_RETURNS = dict(zip(ALL_VARS, [0.06039, 0.00723, 0.03307, 0.04507, 0.03798, 0.01689, 0.00050, 0.00050]))
DEFAULT_VOLS = dict(zip(ALL_VARS, [0.135, 0.042, 0.065, 0.116, 0.178, 0.036, 0.005, 0.012]))
DEFAULT_YIELDS = dict(zip(ASSETS, [0.02040, 0.00607, 0.03816, 0.02597, 0.0, 0.0, 0.00050]))

DEFAULT_CORR = np.array([
    # Eq    DMGov  BOther Prop   Gold   HFs    Cash   Infl
    [1.00, -0.01,  0.72,  0.59,  0.00,  0.50, -0.16,  0.04],
    [-0.01, 1.00,  0.30,  0.16,  0.26, -0.17,  0.04, -0.04],
    [0.72,  0.30,  1.00,  0.53,  0.29,  0.47, -0.08, -0.06],
    [0.59,  0.16,  0.53,  1.00,  0.04,  0.31, -0.18,  0.03],
    [0.00,  0.26,  0.29,  0.04,  1.00,  0.27,  0.08, -0.07],
    [0.50, -0.17,  0.47,  0.31,  0.27,  1.00, -0.02,  0.01],
    [-0.16, 0.04, -0.08, -0.18,  0.08, -0.02,  1.00,  0.05],
    [0.04, -0.04, -0.06,  0.03, -0.07,  0.01,  0.05,  1.00],
])

# Current composite weights (19 indices) as provided by AG.
DEFAULT_MAP = [
    ("Equity", "MSCI UK", 0.062761),
    ("Equity", "MSCI UK Small Cap", 0.010993),
    ("Equity", "MSCI North America", 0.558354),
    ("Equity", "MSCI Japan", 0.045810),
    ("Equity", "MSCI Europe ex UK", 0.101585),
    ("Equity", "MSCI Pacific ex Japan", 0.023100),
    ("Equity", "MSCI Emerging Markets", 0.092946),
    ("Equity", "MSCI AC World ex UK Small Cap", 0.104451),
    ("DMGovt", "ICE BofA UK Conventional Gilts", 0.357612),
    ("DMGovt", "ICE BofA UK Inflation-Linked Gilts", 0.142388),
    ("DMGovt", "ICE BofA US Treasuries", 0.448340),
    ("DMGovt", "ICE BofA US Inflation-Linked Treasuries", 0.051660),
    ("BondOther", "ICE BofA Global Corporate", 0.600000),
    ("BondOther", "ICE BofA Global High Yield", 0.262998),
    ("BondOther", "ICE BofA Emerging Markets External Sovereign", 0.137002),
    ("Property", "S&P Global REIT", 1.0),
    ("Gold", "Gold", 1.0),
    ("HFs", "SG CTA", 1.0),
    ("Cash", "1M SONIA", 1.0),
]

CONSTRAINT_ROWS = [
    "Annualised Volatility",
    "Min Equity", "Max Equity",
    "Min DMGovt", "Max DMGovt",
    "Min BondOther", "Max BondOther",
    "Min Property", "Max Property",
    "Min Gold", "Max Gold",
    "Min HFs", "Max HFs",
    "Min Cash", "Max Cash",
    "Min Yield",
]

# Illustrative strategies only - placeholders for the Bank's actual strategy set.
#                 vol    Eq min/max   DMG min/max  BO min/max   Prop       Gold       HFs        Cash       MinYld
DEFAULT_SETUP = {
    "Strategy_1": [0.05, 0.00, 0.30, 0.00, 0.60, 0.00, 0.40, 0.00, 0.10, 0.00, 0.10, 0.00, 0.15, 0.00, 0.20, 0.000],
    "Strategy_2": [0.07, 0.20, 0.50, 0.00, 0.50, 0.00, 0.40, 0.00, 0.10, 0.00, 0.10, 0.00, 0.15, 0.00, 0.15, 0.000],
    "Strategy_3": [0.09, 0.40, 0.70, 0.00, 0.40, 0.00, 0.35, 0.00, 0.10, 0.00, 0.10, 0.00, 0.15, 0.00, 0.10, 0.000],
    "Strategy_4": [0.11, 0.60, 0.85, 0.00, 0.30, 0.00, 0.30, 0.00, 0.10, 0.00, 0.10, 0.00, 0.10, 0.00, 0.10, 0.000],
    "Strategy_5": [0.13, 0.75, 1.00, 0.00, 0.20, 0.00, 0.20, 0.00, 0.10, 0.00, 0.10, 0.00, 0.10, 0.00, 0.05, 0.000],
    "Strategy_6": [0.09, 0.40, 0.70, 0.00, 0.40, 0.00, 0.45, 0.00, 0.10, 0.00, 0.10, 0.00, 0.15, 0.00, 0.10, 0.025],
}

DEFAULT_SETTINGS = {
    # name: (value, description)
    "HorizonYears": (20, "Projection horizon in years; objective = mean gross value at this year (assumed W97)"),
    "InitialValue": (100, "Starting portfolio value V0. Matters because penalties are in absolute objective units"),
    "Distribution": ("Normal", "Marginal distribution of annual returns: Normal or Lognormal (lognormal on 1+r)"),
    "VolConstraintType": ("max", "'max' = volatility must not exceed target; 'target' = penalise deviation either side"),
    "PenaltyFactorVolYield": (1000, "Penalty = (factor x deviation)^2 for Annualised Volatility and Min Yield"),
    "PenaltyFactorOther": (100, "Penalty = (factor x deviation)^2 for all asset-class min/max constraints"),
    "WeightUnits": (1000, "Integer grid for weights (RISKOptimizer adjustable cells 0..1000)"),
    "Pass1Iterations": (500, "Monte Carlo iterations per trial, optimisation pass 1"),
    "Pass2Iterations": (1000, "Monte Carlo iterations per trial, optimisation pass 2 (seeded from pass 1)"),
    "FinalIterations": (10000, "Iterations for the final simulation of the optimised portfolios"),
    "MaxTrials": (20000, "Max trials (objective evaluations) per pass. Instructions use 1,000 for RISKOptimizer; "
                         "the challenger's search algorithm differs, so counts are not like-for-like"),
    "MinTrials": (10000, "Progress-based stopping is only applied after this many trials"),
    "MaxMinutes": (60, "Maximum run time per optimisation pass, in minutes"),
    "ConvergenceChangePct": (0.001, "Stop if best objective improves by less than this % ..."),
    "ConvergenceTrials": (100, "... over this many trials"),
    "PopulationPerAsset": (5, "Search population size = this x number of assets"),
    "Seed": (12345, "Random seed for reproducibility"),
}


# =====================================================================================
# 2. INPUT HANDLING
# =====================================================================================
@dataclass
class Inputs:
    mu: pd.Series                 # expected returns, ALL_VARS
    vol: pd.Series                # volatilities, ALL_VARS
    yld: pd.Series                # yields, ASSETS
    corr: pd.DataFrame            # rank correlation, ALL_VARS x ALL_VARS
    amap: pd.DataFrame            # Composite, Index, Weight
    setup: pd.DataFrame           # CONSTRAINT_ROWS x strategies
    settings: dict
    source: str
    notes: list = field(default_factory=list)


def default_inputs() -> Inputs:
    return Inputs(
        mu=pd.Series(DEFAULT_RETURNS)[ALL_VARS],
        vol=pd.Series(DEFAULT_VOLS)[ALL_VARS],
        yld=pd.Series(DEFAULT_YIELDS)[ASSETS],
        corr=pd.DataFrame(DEFAULT_CORR, index=ALL_VARS, columns=ALL_VARS),
        amap=pd.DataFrame(DEFAULT_MAP, columns=["Composite", "Index", "Weight"]),
        setup=pd.DataFrame(DEFAULT_SETUP, index=CONSTRAINT_ROWS),
        settings={k: v[0] for k, v in DEFAULT_SETTINGS.items()},
        source="embedded defaults",
    )


def _read_vector(xl: pd.ExcelFile, sheet: str, value_col: str, keys: list[str]) -> pd.Series:
    df = xl.parse(sheet)
    df.columns = [str(c).strip() for c in df.columns]
    if "AssetClass" not in df.columns or value_col not in df.columns:
        raise ValueError(f"Sheet '{sheet}' must have columns 'AssetClass' and '{value_col}'")
    s = df.dropna(subset=["AssetClass"]).set_index(df["AssetClass"].dropna().astype(str).str.strip())[value_col]
    missing = [k for k in keys if k not in s.index]
    if missing:
        raise ValueError(f"Sheet '{sheet}' is missing asset classes: {missing}")
    return s[keys].astype(float)


def load_inputs(path: str | None) -> Inputs:
    if path is None:
        return default_inputs()
    p = Path(path)
    if not p.exists():
        warnings.warn(f"Input file '{path}' not found - using embedded defaults.")
        return default_inputs()

    xl = pd.ExcelFile(p)
    inp = default_inputs()
    inp.source = str(p)

    inp.mu = _read_vector(xl, "ReturnAssumptions", "ExpectedReturn", ALL_VARS)
    inp.vol = _read_vector(xl, "Volatility", "Volatility", ALL_VARS)
    inp.yld = _read_vector(xl, "Yield", "Yield", ASSETS)

    c = xl.parse("CorrelationMatrix", index_col=0)
    c.index = [str(i).strip() for i in c.index]
    c.columns = [str(i).strip() for i in c.columns]
    missing = [k for k in ALL_VARS if k not in c.index or k not in c.columns]
    if missing:
        raise ValueError(f"CorrelationMatrix is missing: {missing}")
    inp.corr = c.loc[ALL_VARS, ALL_VARS].astype(float)

    m = xl.parse("AssetClassMap", usecols="A:C")
    m.columns = ["Composite", "Index", "Weight"]
    inp.amap = m.dropna(subset=["Index"]).assign(
        Composite=lambda d: d["Composite"].ffill().astype(str).str.strip(),
        Weight=lambda d: d["Weight"].astype(float),
    ).reset_index(drop=True)

    s = xl.parse("Setup", index_col=0)
    s.index = [str(i).strip() for i in s.index]
    s = s[[c for c in s.columns if not str(c).startswith("Unnamed") and str(c).strip()]]
    missing = [r for r in CONSTRAINT_ROWS if r not in s.index]
    if missing:
        raise ValueError(f"Setup sheet is missing rows: {missing}")
    inp.setup = s.loc[CONSTRAINT_ROWS].astype(float)

    if "Settings" in xl.sheet_names:
        st = xl.parse("Settings")
        st.columns = [str(c).strip() for c in st.columns]
        for _, row in st.dropna(subset=["Parameter"]).iterrows():
            key = str(row["Parameter"]).strip()
            if key in inp.settings and pd.notna(row["Value"]):
                default = DEFAULT_SETTINGS[key][0]
                val = row["Value"]
                inp.settings[key] = type(default)(val) if not isinstance(default, str) else str(val).strip()
    return inp


def nearest_correlation(a: np.ndarray, iters: int = 200) -> np.ndarray:
    """Higham (2002) alternating projections - nearest valid correlation matrix."""
    y, ds = a.copy(), np.zeros_like(a)
    for _ in range(iters):
        r = y - ds
        vals, vecs = np.linalg.eigh(r)
        x = vecs @ np.diag(np.clip(vals, 1e-10, None)) @ vecs.T
        ds = x - r
        y = x.copy()
        np.fill_diagonal(y, 1.0)
    return (y + y.T) / 2


def validate(inp: Inputs) -> list[str]:
    """Input checks; returns a list of findings (also stored on inp.notes)."""
    notes = []
    c = inp.corr.values
    if not np.allclose(c, c.T, atol=1e-8):
        notes.append("Correlation matrix is not symmetric - symmetrised as (C + C')/2.")
        c = (c + c.T) / 2
    if not np.allclose(np.diag(c), 1.0):
        notes.append("Correlation matrix diagonal is not 1 - diagonal reset to 1.")
        np.fill_diagonal(c, 1.0)
    min_eig = float(np.linalg.eigvalsh(c).min())
    try:
        np.linalg.cholesky(c)
        notes.append(f"Correlation matrix passes Cholesky check (min eigenvalue {min_eig:.4f}).")
    except np.linalg.LinAlgError:
        c = nearest_correlation(c)
        notes.append(f"Correlation matrix FAILS Cholesky (min eigenvalue {min_eig:.4f}) - "
                     "replaced by nearest valid correlation matrix (Higham).")
    inp.corr = pd.DataFrame(c, index=ALL_VARS, columns=ALL_VARS)

    sums = inp.amap.groupby("Composite")["Weight"].sum()
    for comp, tot in sums.items():
        if abs(tot - 1) > 1e-4:
            notes.append(f"AssetClassMap: '{comp}' index weights sum to {tot:.4%}, not 100%.")
    unknown = set(inp.amap["Composite"]) - set(ASSETS)
    if unknown:
        notes.append(f"AssetClassMap: unrecognised composites {sorted(unknown)} (look-through will ignore them).")

    for strat in inp.setup.columns:
        s = inp.setup[strat]
        mins = np.array([s[f"Min {a}"] for a in ASSETS])
        maxs = np.array([s[f"Max {a}"] for a in ASSETS])
        if (mins > maxs + 1e-12).any():
            notes.append(f"{strat}: at least one Min exceeds its Max.")
        if mins.sum() > 1 + 1e-9 or maxs.sum() < 1 - 1e-9:
            notes.append(f"{strat}: bounds infeasible (sum of mins {mins.sum():.1%}, sum of maxes {maxs.sum():.1%}).")
    inp.notes = notes
    return notes


# =====================================================================================
# 3. SIMULATION (Iman-Conover rank correlation, as used by @RISK)
# =====================================================================================
def marginal_draws(n: int, mu: np.ndarray, vol: np.ndarray, dist: str, rng) -> np.ndarray:
    """Independent draws of annual returns for each variable (columns)."""
    z = rng.standard_normal((n, len(mu)))
    if dist.lower() == "lognormal":
        s2 = np.log1p(vol**2 / (1 + mu) ** 2)
        m = np.log1p(mu) - s2 / 2
        return np.exp(m + np.sqrt(s2) * z) - 1
    return mu + vol * z


def iman_conover(x: np.ndarray, target: np.ndarray, rng) -> np.ndarray:
    """Re-order each column of x so that the rank correlation approximates `target`,
    leaving every marginal distribution untouched (Iman & Conover, 1982)."""
    n, k = x.shape
    scores = stats.norm.ppf(np.arange(1, n + 1) / (n + 1))          # van der Waerden scores
    s = np.column_stack([rng.permutation(scores) for _ in range(k)])
    e = np.corrcoef(s, rowvar=False)
    f = np.linalg.cholesky(e)
    p = np.linalg.cholesky(target)
    t = s @ np.linalg.inv(f).T @ p.T                                  # scores with target correlation
    out = np.empty_like(x)
    for j in range(k):
        ranks = stats.rankdata(t[:, j], method="ordinal").astype(int) - 1
        out[:, j] = np.sort(x[:, j])[ranks]
    return out


def simulate(inp: Inputs, n_iter: int, seed: int) -> np.ndarray:
    """Returns array (n_iter, horizon, n_vars) of correlated annual returns."""
    h = int(inp.settings["HorizonYears"])
    rng = np.random.default_rng(seed)
    x = marginal_draws(n_iter * h, inp.mu.values, inp.vol.values, inp.settings["Distribution"], rng)
    x = iman_conover(x, inp.corr.values, rng)
    return x.reshape(n_iter, h, len(ALL_VARS))


# =====================================================================================
# 4. PORTFOLIO METRICS AND PENALISED OBJECTIVE
# =====================================================================================
def portfolio_paths(w: np.ndarray, sims: np.ndarray):
    r = sims[:, :, : len(ASSETS)] @ w                      # (n, H) annual portfolio returns
    return r


def penalties(w: np.ndarray, vol: float, port_yield: float, cons: pd.Series, st: dict) -> dict:
    kv, ko = st["PenaltyFactorVolYield"], st["PenaltyFactorOther"]
    out = {}
    tgt = cons["Annualised Volatility"]
    dev = abs(vol - tgt) if st["VolConstraintType"].lower() == "target" else max(0.0, vol - tgt)
    out["Annualised Volatility"] = (kv * dev) ** 2
    out["Min Yield"] = (kv * max(0.0, cons["Min Yield"] - port_yield)) ** 2
    for i, a in enumerate(ASSETS):
        out[f"Min {a}"] = (ko * max(0.0, cons[f"Min {a}"] - w[i])) ** 2
        out[f"Max {a}"] = (ko * max(0.0, w[i] - cons[f"Max {a}"])) ** 2
    return out


def evaluate(w: np.ndarray, sims: np.ndarray, inp: Inputs, cons: pd.Series) -> dict:
    r = portfolio_paths(w, sims)
    v0 = inp.settings["InitialValue"]
    terminal = v0 * np.prod(1 + r, axis=1)
    vol = float(r.std(ddof=1))
    yld = float(w @ inp.yld.values)
    pen = penalties(w, vol, yld, cons, inp.settings)
    mean_tv = float(terminal.mean())
    return {"mean_terminal": mean_tv, "vol": vol, "yield": yld,
            "penalty": sum(pen.values()), "objective": mean_tv - sum(pen.values()), "penalties": pen}


class TrialTracker:
    """Implements RISKOptimizer-style stopping: max trials, max time, and
    'progress' (improvement < x% over the last N trials)."""

    def __init__(self, st: dict):
        self.max_trials = int(st["MaxTrials"])
        self.max_secs = float(st["MaxMinutes"]) * 60
        self.chg = float(st["ConvergenceChangePct"]) / 100
        self.window = int(st["ConvergenceTrials"])
        self.min_trials = int(st["MinTrials"])
        self.best_hist: list[float] = []
        self.best = -np.inf
        self.t0 = time.time()

    def record(self, val: float):
        self.best = max(self.best, val)
        self.best_hist.append(self.best)

    def should_stop(self) -> bool:
        n = len(self.best_hist)
        if n >= self.max_trials or time.time() - self.t0 > self.max_secs:
            return True
        if n >= self.min_trials and n > self.window:
            old, new = self.best_hist[-self.window - 1], self.best_hist[-1]
            if np.isfinite(old) and abs(new - old) <= self.chg * abs(old):
                return True
        return False


def optimise_replica(inp: Inputs, cons: pd.Series, n_iter: int, seed: int, x0=None):
    st = inp.settings
    units = int(st["WeightUnits"])
    sims = simulate(inp, n_iter, seed)          # common random numbers within a pass
    tracker = TrialTracker(st)

    def neg_obj(x):
        x = np.round(x)
        tot = x.sum()
        if tot <= 0:
            val = -1e12
        else:
            val = evaluate(x / tot, sims, inp, cons)["objective"]
        tracker.record(val)
        return -val

    res = differential_evolution(
        neg_obj,
        bounds=[(0, units)] * len(ASSETS),
        integrality=[True] * len(ASSETS),
        popsize=int(st["PopulationPerAsset"]),
        maxiter=100000,                     # stopping handled by tracker
        tol=0, atol=0,
        mutation=(0.5, 1.0), recombination=0.7,
        seed=seed, polish=False, x0=x0, init="latinhypercube",
        callback=lambda *a, **k: tracker.should_stop(),
        updating="immediate",
    )
    x = np.round(res.x)
    return x, x / x.sum(), len(tracker.best_hist)


def optimise_analytical(inp: Inputs, cons: pd.Series, seed: int):
    mu = inp.mu[ASSETS].values
    sd = inp.vol[ASSETS].values
    cov = np.outer(sd, sd) * inp.corr.loc[ASSETS, ASSETS].values
    y = inp.yld.values
    lo = np.array([cons[f"Min {a}"] for a in ASSETS])
    hi = np.array([cons[f"Max {a}"] for a in ASSETS])
    tgt = cons["Annualised Volatility"]
    vol_type = inp.settings["VolConstraintType"].lower()

    constraints = [{"type": "eq", "fun": lambda w: w.sum() - 1},
                   {"type": "ineq", "fun": lambda w: w @ y - cons["Min Yield"]}]
    if vol_type == "target":
        constraints.append({"type": "eq", "fun": lambda w: tgt**2 - w @ cov @ w})
    else:
        constraints.append({"type": "ineq", "fun": lambda w: tgt**2 - w @ cov @ w})

    rng = np.random.default_rng(seed)
    best = None
    for k in range(30):  # multi-start for robustness
        w0 = rng.dirichlet(np.ones(len(ASSETS))) if k else np.clip(np.full(len(ASSETS), 1 / len(ASSETS)), lo, hi)
        r = minimize(lambda w: -w @ mu, w0, method="SLSQP", bounds=list(zip(lo, hi)),
                     constraints=constraints, options={"ftol": 1e-12, "maxiter": 500})
        if r.success and (best is None or r.fun < best.fun):
            best = r
    if best is None:
        return None, "Infeasible: no portfolio satisfies all hard constraints"
    return np.clip(best.x, 0, None) / np.clip(best.x, 0, None).sum(), "OK"


# =====================================================================================
# 5. FINAL SIMULATION AND REPORTING
# =====================================================================================
def final_stats(w, inp: Inputs, cons: pd.Series, sims: np.ndarray) -> dict:
    ev = evaluate(w, sims, inp, cons)
    r = portfolio_paths(w, sims)
    v0 = inp.settings["InitialValue"]
    terminal = v0 * np.prod(1 + r, axis=1)
    infl = np.prod(1 + sims[:, :, -1], axis=1)
    real = terminal / infl
    sd = inp.vol[ASSETS].values
    cov = np.outer(sd, sd) * inp.corr.loc[ASSETS, ASSETS].values
    breaches = [k for k, v in ev["penalties"].items() if v > 1e-12]
    return {
        "Expected return w'mu": float(w @ inp.mu[ASSETS].values),
        "Analytic volatility sqrt(w'Sw)": float(np.sqrt(w @ cov @ w)),
        "Simulated annual volatility": ev["vol"],
        "Portfolio yield": ev["yield"],
        "Mean terminal value (W97 proxy)": ev["mean_terminal"],
        "Median terminal value": float(np.median(terminal)),
        "5th pct terminal value": float(np.percentile(terminal, 5)),
        "95th pct terminal value": float(np.percentile(terminal, 95)),
        "Mean real terminal value": float(real.mean()),
        "Total penalty": ev["penalty"],
        "Penalised objective": ev["objective"],
        "Constraints breached": ", ".join(breaches) if breaches else "None",
    }


def run(inp: Inputs, strategies: list[str] | None = None, verbose: bool = True):
    notes = validate(inp)
    st = inp.settings
    strategies = strategies or list(inp.setup.columns)
    seed = int(st["Seed"])
    final_sims = simulate(inp, int(st["FinalIterations"]), seed + 999)

    # achieved rank correlation diagnostic
    flat = final_sims.reshape(-1, len(ALL_VARS))
    achieved = stats.spearmanr(flat).correlation
    max_corr_err = float(np.abs(achieved - inp.corr.values).max())
    notes.append(f"Max abs difference achieved vs target rank correlation "
                 f"({st['FinalIterations']} x {st['HorizonYears']} draws): {max_corr_err:.4f}")

    weights_rows, stats_rows, log_rows = [], [], []
    for i, s in enumerate(strategies):
        cons = inp.setup[s]
        t0 = time.time()
        if verbose:
            print(f"[{s}] replica pass 1 ...", flush=True)
        x1, w1, n1 = optimise_replica(inp, cons, int(st["Pass1Iterations"]), seed + 10 * i)
        if verbose:
            print(f"[{s}] replica pass 2 ...", flush=True)
        x2, w2, n2 = optimise_replica(inp, cons, int(st["Pass2Iterations"]), seed + 10 * i + 1, x0=x1)
        wa, status = optimise_analytical(inp, cons, seed)
        log_rows.append({"Strategy": s, "Pass 1 trials": n1, "Pass 2 trials": n2,
                         "Pass 1 raw integers": " ".join(map(str, x1.astype(int))),
                         "Pass 2 raw integers": " ".join(map(str, x2.astype(int))),
                         "Analytical status": status, "Runtime (s)": round(time.time() - t0, 1)})

        for method, w in [("Replica", w2), ("Analytical", wa)]:
            if w is None:
                continue
            weights_rows.append({"Strategy": s, "Method": method, **dict(zip(ASSETS, w))})
            stats_rows.append({"Strategy": s, "Method": method, **final_stats(w, inp, cons, final_sims)})
        if verbose:
            print(f"[{s}] done in {time.time() - t0:.1f}s", flush=True)

    weights = pd.DataFrame(weights_rows)
    summary = pd.DataFrame(stats_rows)

    # look-through to index level
    lt = []
    for _, row in weights.iterrows():
        for _, m in inp.amap.iterrows():
            if m["Composite"] in ASSETS:
                lt.append({"Strategy": row["Strategy"], "Method": row["Method"],
                           "Composite": m["Composite"], "Index": m["Index"],
                           "Weight": row[m["Composite"]] * m["Weight"]})
    lookthrough = (pd.DataFrame(lt)
                   .pivot_table(index=["Composite", "Index"], columns=["Strategy", "Method"],
                                values="Weight", sort=False)
                   if lt else pd.DataFrame())
    return {"weights": weights, "summary": summary, "lookthrough": lookthrough,
            "log": pd.DataFrame(log_rows), "notes": notes}


def write_results(res: dict, inp: Inputs, path: str):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        res["weights"].to_excel(xw, sheet_name="Weights", index=False)
        res["summary"].to_excel(xw, sheet_name="Summary", index=False)
        if not res["lookthrough"].empty:
            res["lookthrough"].to_excel(xw, sheet_name="LookThrough")
        res["log"].to_excel(xw, sheet_name="OptimiserLog", index=False)
        diag = pd.DataFrame({"Diagnostics": [f"Input source: {inp.source}"] + res["notes"]
                             + [f"Setting {k} = {v}" for k, v in inp.settings.items()]})
        diag.to_excel(xw, sheet_name="Diagnostics", index=False)

    wb = load_workbook(path)
    pct_cols = set(ASSETS) | {"Expected return w'mu", "Analytic volatility sqrt(w'Sw)",
                              "Simulated annual volatility", "Portfolio yield"}
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial")
        headers = [c.value for c in ws[1]]
        for j, h in enumerate(headers, start=1):
            fmt = "0.00%" if h in pct_cols else ("#,##0.00" if isinstance(h, str) and "value" in h.lower()
                                                  or h in ("Total penalty", "Penalised objective") else None)
            if fmt:
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    row[0].number_format = fmt
            ws.column_dimensions[get_column_letter(j)].width = max(12, min(45, len(str(h or "")) + 4))
        if ws.title == "LookThrough":
            for row in ws.iter_rows(min_row=4, min_col=3):
                for cell in row:
                    cell.number_format = "0.00%"
            ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 14, 46
        if ws.title == "Diagnostics":
            ws.column_dimensions["A"].width = 120
    wb.save(path)


# =====================================================================================
# 6. TEMPLATE BUILDER
# =====================================================================================
def make_template(path: str):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    inp_font = Font(name="Arial", color="0000FF")
    inp_fill = PatternFill("solid", fgColor="FFFF00")
    base = Font(name="Arial")
    note_font = Font(name="Arial", italic=True, color="808080")
    src_2022 = "Illustrative: 2022 developer documentation (as at 30/11/2021). Replace with current values."

    wb = Workbook()
    ws = wb.active
    ws.title = "ReadMe"
    lines = [
        ("SAA Challenger Optimiser - input template", True),
        ("", False),
        ("How to use", True),
        ("Edit ONLY the cells with blue text on yellow fill. Do not rename tabs, headers or AssetClass codes.", False),
        ("Percentages are stored as decimals (e.g. 0.135 shows as 13.5%).", False),
        ("Asset class codes used throughout: Equity, DMGovt, BondOther, Property, Gold, HFs, Cash, Inflation.", False),
        ("Setup: one column per strategy - add or remove strategy columns freely (header = strategy name).", False),
        ("Rows below the Setup table and check cells on any tab are ignored by the script.", False),
        ("", False),
        ("Tabs", True),
        ("ReturnAssumptions / Volatility / Yield - composite-level inputs (Inflation has no yield).", False),
        ("CorrelationMatrix - 8x8 rank (Spearman) correlation matrix, incl. Inflation.", False),
        ("AssetClassMap - 19 indices and their weights within each composite (used for look-through output).", False),
        ("Setup - constraints per strategy. Settings - optimiser and simulation settings.", False),
        ("", False),
        ("Caveats", True),
        ("Composite returns, volatilities, yields and correlations are illustrative values from the 2022 documentation.", False),
        ("Hedge fund inputs relate to the former HFRX composite, not the current SG CTA index.", False),
        ("Strategy_1 to Strategy_6 constraints are placeholders, not the Bank's actual strategy parameters.", False),
        ("MaxTrials defaults to 20,000 (MinTrials 10,000): with 1,000 trials the challenger's search can stop well short of the optimum.", False),
        ("Run: python saa_challenger.py --input <this file> --output <results file>", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        ws.cell(i, 1, txt).font = Font(name="Arial", bold=bold, size=12 if i == 1 else 10)
    ws.cell(4, 3, "Example input cell").font = inp_font
    ws.cell(4, 3).fill = inp_fill
    ws.column_dimensions["A"].width = 110
    ws.column_dimensions["C"].width = 20

    def header(sheet, cols):
        for j, c in enumerate(cols, start=1):
            cell = sheet.cell(1, j, c)
            cell.font, cell.fill = hdr_font, hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def vector_sheet(name, col, data, keys):
        sh = wb.create_sheet(name)
        header(sh, ["AssetClass", col, "Description", "Source / note"])
        for i, k in enumerate(keys, start=2):
            sh.cell(i, 1, k).font = base
            c = sh.cell(i, 2, data[k])
            c.font, c.fill, c.number_format = inp_font, inp_fill, "0.000%"
            sh.cell(i, 3, ASSET_DESCRIPTIONS[k]).font = base
            sh.cell(i, 4, src_2022).font = note_font
        sh.column_dimensions["A"].width, sh.column_dimensions["B"].width = 14, 16
        sh.column_dimensions["C"].width, sh.column_dimensions["D"].width = 36, 80
        return sh

    vector_sheet("ReturnAssumptions", "ExpectedReturn", DEFAULT_RETURNS, ALL_VARS)
    v = vector_sheet("Volatility", "Volatility", DEFAULT_VOLS, ALL_VARS)
    v.cell(2, 2).comment = Comment("Mapped from the 'Expected Annual Returns & Volatility' chart: the printed "
                                   "composite volatility table in the 2022 document appears shifted by one row.", "Template")
    vector_sheet("Yield", "Yield", DEFAULT_YIELDS, ASSETS)

    sh = wb.create_sheet("CorrelationMatrix")
    header(sh, ["AssetClass"] + ALL_VARS)
    for i, k in enumerate(ALL_VARS, start=2):
        sh.cell(i, 1, k).font = Font(name="Arial", bold=True)
        for j in range(len(ALL_VARS)):
            c = sh.cell(i, j + 2, float(DEFAULT_CORR[i - 2, j]))
            c.font, c.number_format = inp_font, "0.00"
            if j != i - 2:
                c.fill = inp_fill
    r0 = len(ALL_VARS) + 3
    sh.cell(r0, 1, "Check: max |C - C'| (should be 0)").font = note_font
    sh.cell(r0, 4, "=SUMPRODUCT(ABS(B2:I9-TRANSPOSE(B2:I9)))").font = base
    sh.cell(r0 + 1, 1, src_2022).font = note_font
    sh.column_dimensions["A"].width = 14
    for col in "BCDEFGHI":
        sh.column_dimensions[col].width = 11

    sh = wb.create_sheet("AssetClassMap")
    header(sh, ["Composite", "Index", "Weight in Composite"])
    for i, (comp, idx, w) in enumerate(DEFAULT_MAP, start=2):
        sh.cell(i, 1, comp).font = base
        sh.cell(i, 2, idx).font = base
        c = sh.cell(i, 3, w)
        c.font, c.fill, c.number_format = inp_font, inp_fill, "0.0000%"
    last = len(DEFAULT_MAP) + 1
    sh.cell(1, 6, "Check: Composite").font, sh.cell(1, 6).fill = hdr_font, hdr_fill
    sh.cell(1, 7, "Sum of weights").font, sh.cell(1, 7).fill = hdr_font, hdr_fill
    for i, a in enumerate(ASSETS, start=2):
        sh.cell(i, 6, a).font = base
        c = sh.cell(i, 7, f"=SUMIF($A$2:$A${last},F{i},$C$2:$C${last})")
        c.font, c.number_format = base, "0.0000%"
    sh.cell(last + 2, 1, "Source: current composite weights provided by AG (19 indices).").font = note_font
    sh.column_dimensions["A"].width, sh.column_dimensions["B"].width = 14, 48
    sh.column_dimensions["C"].width, sh.column_dimensions["F"].width, sh.column_dimensions["G"].width = 20, 18, 16

    sh = wb.create_sheet("Setup")
    strats = list(DEFAULT_SETUP.keys())
    header(sh, ["Constraint"] + strats)
    for i, row in enumerate(CONSTRAINT_ROWS, start=2):
        sh.cell(i, 1, row).font = base
        for j, s in enumerate(strats, start=2):
            c = sh.cell(i, j, DEFAULT_SETUP[s][i - 2])
            c.font, c.fill, c.number_format = inp_font, inp_fill, "0.0%"
    last = len(CONSTRAINT_ROWS) + 1
    chk = last + 2
    sh.cell(chk, 1, "Check: sum of mins (<= 100%)").font = note_font
    sh.cell(chk + 1, 1, "Check: sum of maxes (>= 100%)").font = note_font
    min_rows = [CONSTRAINT_ROWS.index(f"Min {a}") + 2 for a in ASSETS]
    max_rows = [CONSTRAINT_ROWS.index(f"Max {a}") + 2 for a in ASSETS]
    for j in range(2, len(strats) + 2):
        col = sh.cell(1, j).column_letter
        sh.cell(chk, j, "=" + "+".join(f"{col}{r}" for r in min_rows)).number_format = "0.0%"
        sh.cell(chk + 1, j, "=" + "+".join(f"{col}{r}" for r in max_rows)).number_format = "0.0%"
        sh.cell(chk, j).font = sh.cell(chk + 1, j).font = base
    sh.cell(chk + 3, 1, "Strategy columns are illustrative placeholders. Annualised Volatility is the "
                        "target volatility; Min Yield set to 0% where not applicable.").font = note_font
    sh.column_dimensions["A"].width = 30
    for j in range(2, len(strats) + 2):
        sh.column_dimensions[sh.cell(1, j).column_letter].width = 13

    sh = wb.create_sheet("Settings")
    header(sh, ["Parameter", "Value", "Description"])
    for i, (k, (val, desc)) in enumerate(DEFAULT_SETTINGS.items(), start=2):
        sh.cell(i, 1, k).font = base
        c = sh.cell(i, 2, val)
        c.font, c.fill = inp_font, inp_fill
        sh.cell(i, 3, desc).font = base
    sh.column_dimensions["A"].width, sh.column_dimensions["B"].width, sh.column_dimensions["C"].width = 26, 12, 95

    wb.save(path)


# =====================================================================================
# 7. CLI
# =====================================================================================
def main(argv=None):
    ap = argparse.ArgumentParser(description="SAA challenger optimiser (validation tool)")
    ap.add_argument("--input", help="Excel input file (template format). Omit to use embedded defaults.")
    ap.add_argument("--output", default="SAA_Challenger_Results.xlsx", help="Results workbook")
    ap.add_argument("--strategies", nargs="*", help="Subset of strategy names from the Setup tab")
    ap.add_argument("--make-template", metavar="PATH", help="Write the input template and exit")
    a = ap.parse_args(argv)

    if a.make_template:
        make_template(a.make_template)
        print(f"Template written to {a.make_template}")
        return

    inp = load_inputs(a.input)
    print(f"Inputs loaded from: {inp.source}")
    res = run(inp, a.strategies)
    for n in res["notes"]:
        print(" -", n)
    write_results(res, inp, a.output)
    pd.set_option("display.width", 200)
    print(res["weights"].round(4).to_string(index=False))
    print(f"Results written to {a.output}")


if __name__ == "__main__":
    main()
